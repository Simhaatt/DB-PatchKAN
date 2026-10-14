"""Structural, formula, and cached-value checks for the final workbook."""

from __future__ import annotations

import argparse
import json
import math
from pathlib import Path

from openpyxl import load_workbook

from build_workbook import SHEETS


FORMULA_ERRORS = {
    "#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!",
    "#SPILL!", "#CALC!", "#FIELD!", "#CONNECT!", "#BLOCKED!",
}


def _finite_number(value: object) -> bool:
    return isinstance(value, (int, float)) and math.isfinite(float(value))


def _header_map(ws, row: int) -> dict[str, int]:
    return {
        str(cell.value): cell.column
        for cell in ws[row]
        if cell.value is not None
    }


def verify(path: Path) -> dict[str, object]:
    formula_book = load_workbook(path, data_only=False, read_only=False)
    cached_book = load_workbook(path, data_only=True, read_only=False)
    if formula_book.sheetnames != SHEETS:
        raise AssertionError(f"Unexpected sheet order: {formula_book.sheetnames}")

    formula_count = 0
    cached_errors: list[str] = []
    nonempty: dict[str, int] = {}
    for ws in formula_book.worksheets:
        nonempty[ws.title] = sum(
            1 for row in ws.iter_rows() for cell in row if cell.value is not None
        )
        if nonempty[ws.title] == 0:
            raise AssertionError(f"Empty worksheet: {ws.title}")
        if not ws.print_area:
            raise AssertionError(f"Worksheet has no verification print area: {ws.title}")
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    formula_count += 1
    for ws in cached_book.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value in FORMULA_ERRORS:
                    cached_errors.append(f"{ws.title}!{cell.coordinate}:{cell.value}")
    if cached_errors:
        raise AssertionError(f"Cached formula errors: {cached_errors[:20]}")
    if formula_count == 0:
        raise AssertionError("Workbook contains no formulas")

    parameter_formula_ws = formula_book["01_PARAMETERS"]
    parameter_cached_ws = cached_book["01_PARAMETERS"]
    for row in range(3, 9):
        formula = parameter_formula_ws.cell(row, 9).value
        cached = parameter_cached_ws.cell(row, 9).value
        if not isinstance(formula, str) or not formula.startswith("=$C$"):
            raise AssertionError(f"D-BSI weight I{row} is not linked to authoritative column C")
        if not _finite_number(cached):
            raise AssertionError(f"D-BSI weight I{row} has no recalculated numeric cache")
    weight_total = parameter_cached_ws["I10"].value
    if not _finite_number(weight_total) or abs(float(weight_total) - 1.0) > 1e-12:
        raise AssertionError(f"D-BSI weights do not sum to 1.0: {weight_total!r}")

    bsi_formula_ws = formula_book["11_BSI_RESULTS"]
    bsi_headers = _header_map(bsi_formula_ws, 4)
    required_bsi_columns = {
        "bsi_rate", "bsi_thermal", "bsi_soc", "bsi_cycle", "bsi_transient",
        "bsi_interaction", "bsi_intensity", "stress_dose",
    }
    missing_bsi_columns = sorted(required_bsi_columns.difference(bsi_headers))
    if missing_bsi_columns:
        raise AssertionError(f"D-BSI result sheet is missing columns: {missing_bsi_columns}")
    expected_audit_rows = max(0, bsi_formula_ws.max_row - 4)
    if expected_audit_rows == 0:
        raise AssertionError("D-BSI result sheet contains no records")
    for row in range(5, 5 + expected_audit_rows):
        for column in required_bsi_columns:
            value = cached_book["11_BSI_RESULTS"].cell(row, bsi_headers[column]).value
            if not _finite_number(value):
                raise AssertionError(
                    f"D-BSI result {column} is blank/non-numeric at row {row}; workbook may be incomplete"
                )

    audit_formula_ws = formula_book["18_BSI_AUDIT"]
    if "01_PARAMETERS" not in str(audit_formula_ws["I5"].value):
        raise AssertionError("D-BSI audit contribution is not linked to parameter weights")
    if audit_formula_ws.max_row - 4 != expected_audit_rows:
        raise AssertionError(
            "D-BSI audit/result row mismatch: "
            f"audit={audit_formula_ws.max_row - 4}, results={expected_audit_rows}"
        )

    cached_audit = cached_book["18_BSI_AUDIT"]
    dbsi_diffs: list[float] = []
    dose_diffs: list[float] = []
    statuses: list[str] = []
    required_formula_columns = [*range(9, 16), 17, 18, 20, 21]
    required_numeric_columns = [*range(2, 21)]
    for row in range(5, cached_audit.max_row + 1):
        for column in required_formula_columns:
            if formula_book["18_BSI_AUDIT"].cell(row, column).data_type != "f":
                raise AssertionError(
                    f"Expected audit formula missing at 18_BSI_AUDIT!"
                    f"{formula_book['18_BSI_AUDIT'].cell(row, column).coordinate}"
                )
        for column in required_numeric_columns:
            value = cached_audit.cell(row, column).value
            if not _finite_number(value):
                raise AssertionError(
                    f"Audit cache is blank/non-numeric at 18_BSI_AUDIT!"
                    f"{cached_audit.cell(row, column).coordinate}; run Excel recalculation"
                )
        for column, destination in ((17, dbsi_diffs), (20, dose_diffs)):
            value = cached_audit.cell(row, column).value
            destination.append(abs(float(value)))
        status = cached_audit.cell(row, 21).value
        statuses.append(str(status))
    if len(dbsi_diffs) != expected_audit_rows or len(statuses) != expected_audit_rows:
        raise AssertionError("D-BSI audit cache is incomplete")
    max_dbsi_difference = max(dbsi_diffs)
    max_dose_difference = max(dose_diffs)
    if max_dbsi_difference > 1e-6 or max_dose_difference > 1e-6:
        raise AssertionError(
            f"D-BSI audit mismatch: intensity={max_dbsi_difference}, dose={max_dose_difference}"
        )
    if any(value != "PASS" for value in statuses):
        raise AssertionError("One or more D-BSI audit rows did not pass")

    dashboard_formula_ws = formula_book["17_DASHBOARD"]
    dashboard_cached_ws = cached_book["17_DASHBOARD"]
    dashboard_trip_rows = min(
        127,
        max(0, formula_book["07_TRIP_FEATURES"].max_row - 4),
    )
    for row in range(14, 14 + dashboard_trip_rows):
        for column in range(1, 4):
            formula = dashboard_formula_ws.cell(row, column).value
            if not isinstance(formula, str) or not formula.startswith(
                "='07_TRIP_FEATURES'!"
            ):
                raise AssertionError(
                    "Dashboard trip/chart source was overwritten at "
                    f"17_DASHBOARD!{dashboard_formula_ws.cell(row, column).coordinate}"
                )
            if dashboard_cached_ws.cell(row, column).value is None:
                raise AssertionError(
                    "Dashboard trip/chart cache is blank at "
                    f"17_DASHBOARD!{dashboard_cached_ws.cell(row, column).coordinate}"
                )

    result = {
        "path": str(path.resolve()),
        "sheets": len(formula_book.sheetnames),
        "formula_count": formula_count,
        "cached_formula_errors": 0,
        "max_dbsi_difference": max_dbsi_difference,
        "max_dose_difference": max_dose_difference,
        "audit_rows": expected_audit_rows,
        "nonempty_cells_by_sheet": nonempty,
    }
    formula_book.close()
    cached_book.close()
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    result = verify(args.workbook.resolve())
    text = json.dumps(result, indent=2)
    print(text)
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(text, encoding="utf-8")


if __name__ == "__main__":
    main()
