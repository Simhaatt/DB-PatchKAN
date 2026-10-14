"""Build the thesis-ready DB-PatchKAN Excel research workbook.

The preferred bundled artifact-tool runtime is not available in this workspace
installation, so this deterministic fallback uses openpyxl and is subsequently
recalculated/rendered by the installed Microsoft Excel application.  All D-BSI
audit results are Excel formulas rather than pasted values.
"""

from __future__ import annotations

import argparse
import hashlib
import importlib.metadata
import json
import math
import os
import re
import sys
from pathlib import Path
from typing import Any, Iterable

import numpy as np
import pandas as pd
import yaml
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


SHEETS = [
    "00_README",
    "01_PARAMETERS",
    "02_DATA_DICTIONARY",
    "03_SOURCE_PROFILE",
    "04_QC_SUMMARY",
    "05_RAW_SAMPLE",
    "06_WINDOW_FEATURES",
    "07_TRIP_FEATURES",
    "08_HEALTH_WINDOWS",
    "09_FEATURE_SELECTION",
    "10_BEHAVIOR_CLUSTERS",
    "11_BSI_RESULTS",
    "12_SOH_PREDICTIONS",
    "13_DL_EXPERIMENTS",
    "14_BENCHMARKS_ABLATIONS",
    "15_EXPLAINABILITY",
    "16_EXTERNAL_VALIDATION",
    "17_DASHBOARD",
    "18_BSI_AUDIT",
]

NAVY = "17324D"
TEAL = "0F766E"
TEAL_LIGHT = "DDF3EF"
BLUE = "2563EB"
AMBER = "F59E0B"
RED = "DC2626"
GREEN = "16803A"
SLATE = "475569"
LIGHT = "F4F7FA"
WHITE = "FFFFFF"
GRID = "D8E1E8"

FEATURE_SELECTION_COLUMNS = [
    "Feature_Name", "Feature_Group", "Source_Columns", "Protected_Feature",
    "Missing_Pct", "Dominant_Value_Pct", "Correlation_Group",
    "Representative_Feature", "mRMR_Rank", "ElasticNet_Selection_Frequency",
    "Permutation_Importance", "SHAP_Importance", "EBM_Importance",
    "STG_Gate_Probability", "Channel_Occlusion_Delta",
    "Fold_Selection_Frequency", "Positive_Permutation_Frequency",
    "Positive_Ablation_Frequency", "Stable_Response_Frequency",
    "Stable_Response", "Robustly_Important", "Consensus_Rank", "Final_Status",
    "Rejection_Reason", "Interpretation",
]

DL_EXPERIMENT_COLUMNS = [
    "Run_ID", "Timestamp_UTC", "Environment", "GPU", "Dataset", "Outer_Fold",
    "Inner_Fold", "Seed", "Model_Variant", "Feature_Set", "Config_Hash",
    "Patch_Length", "Patch_Stride", "Embedding_Dim", "Blocks", "Heads",
    "KAN_Grid", "KAN_Shape", "Batch_Size", "Precision", "Best_Epoch",
    "Parameter_Count", "FLOPs", "Train_Time_Min", "Peak_VRAM_GB", "MAE_pp",
    "RMSE_pp", "R2", "Macro_F1", "Interval_Coverage", "Checkpoint_Path",
    "Status", "Notes",
]

SHEET_SPECS = {
    "02_DATA_DICTIONARY": ("Data dictionary", "Every source and derived field, unit, lineage, role, and leakage treatment."),
    "03_SOURCE_PROFILE": ("Source profile", "Counts and robust distribution audit from the cleaned, deterministic source sample."),
    "04_QC_SUMMARY": ("Quality-control summary", "Sentinel/range rules, affected observations, severity, and treatment."),
    "05_RAW_SAMPLE": ("Deterministic raw/clean sample", "Compact, evenly spaced coverage of every trip plus QC-flagged examples; the immutable full source remains external."),
    "06_WINDOW_FEATURES": ("Temporal window features", "One record per valid 256-second segment; overlapping windows stay with their trip."),
    "07_TRIP_FEATURES": ("Trip features", "One record per Nissan trip, including behaviour, operational, target, and D-BSI audit fields."),
    "08_HEALTH_WINDOWS": ("Health and forecast windows", "One record per current/future health checkpoint with chronological lineage."),
    "09_FEATURE_SELECTION": ("Fold-safe feature selection", "Missingness, redundancy, stability, held-out importance, and final decision evidence."),
    "10_BEHAVIOR_CLUSTERS": ("Driver-behaviour discovery", "Fold-local GMM/PCA archetypes, stability evidence, aggression score, and behaviour class."),
    "11_BSI_RESULTS": ("D-BSI results", "Six components, weighted contributions, stress dose, and cumulative dose per trip."),
    "12_SOH_PREDICTIONS": ("SOH predictions", "Targets, predictions, intervals, residuals, model, fold, and split lineage."),
    "13_DL_EXPERIMENTS": ("Deep-learning experiment ledger", "Architecture, environment, hyperparameters, compute, efficiency, metrics, and checkpoint status."),
    "14_BENCHMARKS_ABLATIONS": ("Benchmarks and ablations", "Classical/deep comparisons and the pre-registered ablation matrix."),
    "15_EXPLAINABILITY": ("Explainability evidence", "EBM/SHAP/KAN/gradient/occlusion outputs. Attention alone is not importance."),
    "16_EXTERNAL_VALIDATION": ("External validation", "Dataset mapping, label definition, adapter status, vehicle-wise results, and claim boundary."),
}


def _safe_value(value: Any) -> Any:
    if value is None or value is pd.NA:
        return None
    if isinstance(value, (np.bool_, bool)):
        return bool(value)
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (np.floating, float)):
        value = float(value)
        return value if math.isfinite(value) else None
    if isinstance(value, pd.Timestamp):
        timestamp = value
        if timestamp.tzinfo is not None:
            timestamp = timestamp.tz_convert("UTC").tz_localize(None)
        return timestamp.to_pydatetime()
    if isinstance(value, (dict, list, tuple, set)):
        return json.dumps(value, default=str, sort_keys=True)
    text = str(value)
    return text[:32767]


def _read_csv(path: Path, *, columns: list[str] | None = None) -> pd.DataFrame:
    if not path.exists() or path.stat().st_size == 0:
        return pd.DataFrame(columns=columns or [])
    try:
        frame = pd.read_csv(path, low_memory=False)
    except pd.errors.EmptyDataError:
        frame = pd.DataFrame()
    if columns:
        for column in columns:
            if column not in frame:
                frame[column] = np.nan
        frame = frame[columns + [c for c in frame.columns if c not in columns]]
    # CSV is the interchange format for the research ledgers, but timestamps
    # must remain sortable Excel values rather than opaque strings.  Excel does
    # not accept timezone-aware datetimes, so canonical UTC values are stored as
    # timezone-naive datetimes and labelled explicitly by their column names.
    for column in frame.columns:
        if str(column).casefold().endswith("_utc"):
            parsed = pd.to_datetime(frame[column], errors="coerce", utc=True)
            if parsed.notna().any():
                frame[column] = parsed.dt.tz_convert("UTC").dt.tz_localize(None)
    return frame


def _sha256(path: Path) -> str:
    if not path.exists() or not path.is_file():
        return "pending"
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _package_version(name: str) -> str:
    try:
        return importlib.metadata.version(name)
    except importlib.metadata.PackageNotFoundError:
        return "not-installed"


def _prioritize_columns(frame: pd.DataFrame, priority: Iterable[str]) -> pd.DataFrame:
    ordered = [column for column in priority if column in frame.columns]
    ordered.extend(column for column in frame.columns if column not in ordered)
    return frame.loc[:, ordered]


def _stratified_raw_excerpt(
    frame: pd.DataFrame,
    *,
    regular_per_trip: int = 32,
    anomaly_per_trip: int = 12,
) -> pd.DataFrame:
    """Return a compact, deterministic sample with every trip represented.

    The standalone workbook is an audit/reporting artifact rather than a
    replacement for the immutable CSV/Parquet source.  Keeping hundreds of
    telemetry rows per trip made Excel expand a single worksheet XML beyond
    100 MB.  This excerpt retains evenly spaced coverage for every trip and a
    separate evenly spaced sample of QC-flagged records from every trip.
    """

    if frame.empty or "Trip" not in frame:
        return frame

    sort_columns = [
        column
        for column in ("datetime_utc", "source_row_id", "index")
        if column in frame
    ]
    pieces: list[pd.DataFrame] = []
    for _, trip in frame.groupby("Trip", sort=True, dropna=False):
        ordered = trip.sort_values(sort_columns, kind="stable") if sort_columns else trip

        def evenly_spaced(part: pd.DataFrame, limit: int) -> pd.DataFrame:
            if len(part) <= limit:
                return part
            indices = np.linspace(0, len(part) - 1, num=limit, dtype=int)
            return part.iloc[np.unique(indices)]

        pieces.append(evenly_spaced(ordered, regular_per_trip))
        if "qc_issue_count" in ordered:
            anomalies = ordered[
                pd.to_numeric(ordered["qc_issue_count"], errors="coerce").fillna(0) > 0
            ]
            if not anomalies.empty:
                pieces.append(evenly_spaced(anomalies, anomaly_per_trip))

    excerpt = pd.concat(pieces, ignore_index=False)
    dedupe_columns = [
        column for column in ("source_row_id", "index", "datetime_utc", "Trip")
        if column in excerpt
    ]
    if dedupe_columns:
        excerpt = excerpt.drop_duplicates(subset=dedupe_columns, keep="first")
    if sort_columns:
        excerpt = excerpt.sort_values(sort_columns, kind="stable")
    return excerpt.reset_index(drop=True)


def _first_existing(directory: Path, names: Iterable[str]) -> Path | None:
    for name in names:
        candidate = directory / name
        if candidate.exists():
            return candidate
    return None


def _load_frame(directory: Path, names: Iterable[str], *, columns: list[str] | None = None) -> pd.DataFrame:
    path = _first_existing(directory, names)
    return _read_csv(path, columns=columns) if path else pd.DataFrame(columns=columns or [])


def _table_name(sheet_name: str) -> str:
    return "T_" + re.sub(r"[^A-Za-z0-9_]", "_", sheet_name)


def _style_title(ws, last_col: int, title: str, subtitle: str) -> None:
    last_col = max(4, last_col)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws.cell(1, 1, title)
    ws.cell(1, 1).fill = PatternFill("solid", fgColor=NAVY)
    ws.cell(1, 1).font = Font(color=WHITE, size=18, bold=True)
    ws.cell(1, 1).alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 30
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    ws.cell(2, 1, subtitle)
    ws.cell(2, 1).font = Font(color=SLATE, size=10, italic=True)
    ws.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 29
    ws.sheet_view.showGridLines = False


def _format_header(cells) -> None:
    for cell in cells:
        cell.fill = PatternFill("solid", fgColor=TEAL)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="thin", color=GRID))


def _set_widths(ws, frame: pd.DataFrame, start_col: int = 1) -> None:
    for offset, column in enumerate(frame.columns, start_col):
        samples = [str(column)]
        if not frame.empty:
            samples.extend(str(x) for x in frame[column].dropna().head(100).tolist())
        width = min(42, max(11, max(len(x) for x in samples[:101]) + 2))
        lowered = str(column).lower()
        if "notes" in lowered or "interpret" in lowered or "reason" in lowered or "source_columns" in lowered:
            width = 34
        ws.column_dimensions[get_column_letter(offset)].width = width


def _number_format(column: str) -> str:
    name = column.casefold()
    if name.endswith("_utc") or "timestamp_utc" in name:
        return "yyyy-mm-dd hh:mm:ss"
    if any(token in name for token in ("pct", "frequency", "probability", "coverage")):
        return "0.000"
    if any(token in name for token in ("mae", "rmse", "importance", "delta", "r2", "score", "bsi", "soh", "dose")):
        return "0.000"
    if any(token in name for token in ("count", "rows", "fold", "seed", "rank", "epoch", "blocks", "heads")):
        return "0"
    return "General"


def _write_table_block(
    ws,
    frame: pd.DataFrame,
    *,
    start_row: int,
    table_name: str,
    max_rows: int = 30_000,
) -> tuple[pd.DataFrame, int]:
    """Write a table without changing the surrounding sheet layout."""

    data = frame.head(max_rows).copy()
    if data.shape[1] == 0:
        data = pd.DataFrame({"Status": ["Pending / no records produced"]})
    for col_index, column in enumerate(data.columns, 1):
        ws.cell(start_row, col_index, str(column))
    _format_header(ws[start_row][: data.shape[1]])
    for row_index, row in enumerate(data.itertuples(index=False, name=None), start_row + 1):
        for col_index, value in enumerate(row, 1):
            cell = ws.cell(row_index, col_index, _safe_value(value))
            cell.alignment = Alignment(vertical="top", wrap_text=False)
            cell.number_format = _number_format(str(data.columns[col_index - 1]))
        if row_index % 2:
            for cell in ws[row_index][: data.shape[1]]:
                cell.fill = PatternFill("solid", fgColor="F8FAFC")
    end_row = max(start_row + 1, start_row + len(data))
    if data.empty:
        ws.cell(start_row + 1, 1, "Pending / no records produced")
    ref = f"A{start_row}:{get_column_letter(data.shape[1])}{end_row}"
    table = Table(displayName=table_name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)
    _set_widths(ws, data)
    return data, end_row


def write_frame(
    ws,
    frame: pd.DataFrame,
    title: str,
    subtitle: str,
    *,
    start_row: int = 4,
    max_rows: int = 30_000,
) -> None:
    title_columns = max(1, frame.shape[1])
    _style_title(ws, title_columns, title, subtitle)
    data, end_row = _write_table_block(
        ws,
        frame,
        start_row=start_row,
        table_name=_table_name(ws.title),
        max_rows=max_rows,
    )
    ws.freeze_panes = f"A{start_row + 1}"
    ws.print_title_rows = f"1:{start_row}"
    ws.print_area = f"A1:{get_column_letter(min(data.shape[1], 14))}{min(end_row, 40)}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _flatten(value: Any, prefix: str = "") -> list[tuple[str, Any]]:
    rows: list[tuple[str, Any]] = []
    if isinstance(value, dict):
        for key, child in value.items():
            path = f"{prefix}.{key}" if prefix else str(key)
            rows.extend(_flatten(child, path))
    elif isinstance(value, list):
        rows.append((prefix, json.dumps(value)))
    else:
        rows.append((prefix, value))
    return rows


def build_readme(
    ws,
    manifest: dict[str, Any],
    source_path: Path,
    output_dir: Path,
    config_path: Path,
) -> None:
    _style_title(ws, 8, "DB-PatchKAN Research Workbook", "Driver–battery behaviour, explainable stress, fold-safe feature discovery, and SOH modelling.")
    ws["A4"] = "Evidence status"
    ws["B4"] = "IMPLEMENTED / EXTERNAL VEHICLE VALIDATION PENDING"
    ws["B4"].fill = PatternFill("solid", fgColor="FFF3CD")
    ws["B4"].font = Font(color="7A4B00", bold=True)
    manifest_path = output_dir / "manifest.json"
    software_versions = "; ".join(
        [
            f"Python {sys.version_info.major}.{sys.version_info.minor}.{sys.version_info.micro}",
            f"dbpatchkan {_package_version('dbpatchkan')}",
            f"pandas {_package_version('pandas')}",
            f"numpy {_package_version('numpy')}",
            f"openpyxl {_package_version('openpyxl')}",
        ]
    )
    facts = [
        ("Source path", str(source_path)),
        ("Source SHA-256", manifest.get("source_sha256", "pending")),
        ("Configuration path", str(config_path)),
        ("Configuration SHA-256", _sha256(config_path)),
        ("Manifest path", str(manifest_path)),
        ("Manifest SHA-256", _sha256(manifest_path)),
        ("Rows processed", manifest.get("rows_processed", "pending")),
        ("Pipeline version", manifest.get("pipeline", "pending")),
        ("Pipeline patch", manifest.get("pipeline_patch_level", "pending")),
        ("QC patch", manifest.get("qc_patch_level", "pending")),
        ("D-BSI version", "D-BSI-v1.0"),
        ("Software versions", software_versions),
        ("Vehicle scope", "One longitudinal Nissan Leaf; no row-independent claims"),
        ("SOH definition", "Nissan BMS-reported SOH; kept separate from capacity-based external SOH"),
        ("Output directory", str(output_dir)),
    ]
    ws["A6"] = "Reproducibility fact"
    ws["B6"] = "Value"
    _format_header(ws[6][:2])
    for row, (key, value) in enumerate(facts, 7):
        ws.cell(row, 1, key)
        ws.cell(row, 2, _safe_value(value))

    objectives = [
        "Quantify driver-behaviour influence on immediate battery stress.",
        "Discover stable Eco/Balanced/Aggressive behaviour archetypes.",
        "Reproduce the six-component D-BSI and cumulative stress dose.",
        "Predict current and future SOH through leakage-safe named concepts.",
        "Validate by chronological trip and, when supplied, held-out external vehicle.",
    ]
    ws["A23"] = "Research objectives"
    ws["A23"].font = Font(color=NAVY, bold=True, size=13)
    for row, objective in enumerate(objectives, 24):
        ws.cell(row, 1, row - 23)
        ws.cell(row, 2, objective)

    workflow = [
        ("1", "QC + canonical signals", "Sentinel cleaning, sign audit, native-rate integration, 1 Hz adapter"),
        ("2", "Fold-safe reduction", "Semantic leakage removal → redundancy → mRMR/stability/permutation"),
        ("3", "Behaviour discovery", "Robust scaling → PCA(95%) → stable GMM; supervised follow-up"),
        ("4", "Explainable D-BSI", "Six bounded components, weighted contributions, dose, sensitivity audit"),
        ("5", "SOH modelling", "Classical baselines → Patch Transformer+MLP → DB-PatchKAN"),
        ("6", "Validation", "Expanding time splits, forecast purge, external leave-one-vehicle-out"),
    ]
    ws["D6"] = "Stage"
    ws["E6"] = "Module"
    ws["F6"] = "What is done"
    _format_header(ws[6][3:6])
    for row, values in enumerate(workflow, 7):
        for col, value in enumerate(values, 4):
            ws.cell(row, col, value)
    ws["D15"] = "Interpretation boundary"
    ws["D15"].font = Font(color=NAVY, bold=True, size=13)
    notes = [
        "D-BSI measures driving-related operational stress; it is not a complete charging/calendar-aging index.",
        "Associations are observational and must not be described as causal without a causal design.",
        "The 1.9M telemetry rows do not create 1.9M independent SOH labels.",
        "DB-PatchKAN is recommended only if paired external results beat XGBoost/EBM and a parameter-matched MLP.",
        "A systematic literature and patent review is required before any novelty-first claim.",
    ]
    for row, note in enumerate(notes, 16):
        ws.cell(row, 4, "•")
        ws.cell(row, 5, note)
        ws.merge_cells(start_row=row, start_column=5, end_row=row, end_column=8)

    for col, width in {"A": 22, "B": 54, "C": 3, "D": 13, "E": 26, "F": 34, "G": 16, "H": 16}.items():
        ws.column_dimensions[col].width = width
    for row in range(1, 48):
        ws.row_dimensions[row].height = 22
    ws.freeze_panes = "A6"
    ws.print_area = "A1:H41"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1


def build_parameters(ws, config: dict[str, Any]) -> None:
    _style_title(ws, 10, "Parameters and run configuration", "Column C is the authoritative editable configuration. The D-BSI audit block at right links directly to those weight cells.")
    ws.unmerge_cells("A2:J2")
    ws.merge_cells("A2:F2")
    rows = _flatten({k: v for k, v in config.items() if k != "_meta"})
    parameter_frame = pd.DataFrame(rows, columns=["Parameter_Path", "Value"])
    parameter_frame["Section"] = parameter_frame["Parameter_Path"].str.split(".").str[0]
    parameter_frame["Purpose"] = parameter_frame["Parameter_Path"].map(parameter_purpose)
    parameter_frame = parameter_frame[["Section", "Parameter_Path", "Value", "Purpose"]]
    for col_index, column in enumerate(parameter_frame.columns, 1):
        ws.cell(4, col_index, column)
    _format_header(ws[4][:4])
    for row_index, row in enumerate(parameter_frame.itertuples(index=False, name=None), 5):
        for col_index, value in enumerate(row, 1):
            ws.cell(row_index, col_index, _safe_value(value))
    parameter_rows = {
        str(parameter_frame.iloc[index]["Parameter_Path"]): 5 + index
        for index in range(len(parameter_frame))
    }
    end_row = 4 + len(parameter_frame)
    table = Table(displayName="T_Parameters", ref=f"A4:D{end_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)

    ws["H2"] = "D-BSI component"
    ws["I2"] = "Linked weight"
    ws["J2"] = "Meaning"
    _format_header(ws[2][7:10])
    components = [
        ("Rate", "dbsi.weights.rate", "Electrical discharge/regen C-rate"),
        ("Thermal", "dbsi.weights.thermal", "Pack heat, cold regen, and spread"),
        ("SOC", "dbsi.weights.soc", "High- and low-SOC exposure"),
        ("Cycle", "dbsi.weights.cycle", "Depth of discharge and EFC rate"),
        ("Transient", "dbsi.weights.transient", "C-rate slew / rapid demand changes"),
        ("Interaction", "dbsi.weights.interaction", "Hot discharge and boundary regen"),
    ]
    missing_paths = [path for _, path, _ in components if path not in parameter_rows]
    if missing_paths:
        raise ValueError(f"Missing authoritative D-BSI configuration weights: {missing_paths}")
    for row, (label, path, meaning) in enumerate(components, 3):
        ws.cell(row, 8, label)
        ws.cell(row, 9, f"=$C${parameter_rows[path]}")
        ws.cell(row, 10, f"{meaning} [{path}]")
    ws["H10"] = "Weight check"
    ws["I10"] = "=SUM(I3:I8)"
    ws["J10"] = "Must equal 1.000"
    ws["I3:I10"][0][0].number_format = "0.00%"
    for row in range(3, 11):
        ws.cell(row, 9).number_format = "0.00%"
    ws["I10"].fill = PatternFill("solid", fgColor=TEAL_LIGHT)
    ws.freeze_panes = "A5"
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 42
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 44
    ws.column_dimensions["H"].width = 18
    ws.column_dimensions["I"].width = 12
    ws.column_dimensions["J"].width = 38
    ws.print_area = f"A1:J{min(end_row, 45)}"


def parameter_purpose(path: str) -> str:
    rules = {
        "feature_selection": "Fold-local dimensionality reduction / stability decision",
        "dbsi": "Transparent D-BSI v1 weight or physical threshold",
        "model": "DB-PatchKAN architecture or channel contract",
        "training": "GPU training, mixed precision, optimiser, or reproducibility",
        "validation": "Chronological grouping, purge, horizon, or untouched test rule",
        "clustering": "Behaviour archetype search and bootstrap stability",
        "qc": "Sentinel/range quality-control threshold",
        "data": "Source identity, sign, cadence, integration, or window rule",
        "project": "Run identity and artifact location",
    }
    return rules.get(path.split(".", 1)[0], "Project configuration")


def _feature_summary(frame: pd.DataFrame, limit: int = 12) -> pd.DataFrame:
    if frame.empty or "Feature_Name" not in frame:
        return pd.DataFrame({"Status": ["No consensus-selected features available"]})
    selected = frame.copy()
    if "Stage" in selected and selected["Stage"].astype(str).eq("development_consensus").any():
        selected = selected.loc[selected["Stage"].astype(str).eq("development_consensus")]
    elif "Fold" in selected and selected["Fold"].astype(str).eq("development_consensus").any():
        selected = selected.loc[selected["Fold"].astype(str).eq("development_consensus")]
    if "Final_Status" in selected:
        selected = selected.loc[selected["Final_Status"].astype(str).str.casefold().eq("selected")]
    if "Permutation_Importance" in selected:
        selected["Permutation_Importance"] = pd.to_numeric(
            selected["Permutation_Importance"], errors="coerce"
        )
        selected = selected.sort_values("Permutation_Importance", ascending=False)
    duplicate_keys = [column for column in ("Task", "Feature_Name") if column in selected]
    if duplicate_keys:
        selected = selected.drop_duplicates(duplicate_keys)
    columns = [
        "Task", "Feature_Name", "Feature_Group", "Permutation_Importance",
        "Fold_Selection_Frequency", "Positive_Permutation_Frequency",
        "Positive_Ablation_Frequency", "Stable_Response_Frequency",
        "Robustly_Important", "Consensus_Rank", "Final_Status", "Interpretation",
    ]
    columns = [column for column in columns if column in selected]
    if selected.empty or not columns:
        return pd.DataFrame({"Status": ["No consensus-selected features available"]})
    return selected.loc[:, columns].head(limit).reset_index(drop=True)


def build_feature_selection_sheet(ws, frame: pd.DataFrame) -> pd.DataFrame:
    summary = _feature_summary(frame)
    _style_title(
        ws,
        summary.shape[1],
        "Fold-safe feature selection",
        "Visible summary contains only development-consensus selected features; the complete fold-level audit follows below.",
    )
    ws["A4"] = "Consensus candidates and four-part robustness audit"
    ws["A4"].font = Font(color=NAVY, bold=True, size=12)
    _, summary_end = _write_table_block(
        ws, summary, start_row=5, table_name="T_FeatureSummary", max_rows=12
    )
    detail_start = max(22, summary_end + 4)
    ws.cell(detail_start - 1, 1, "Complete fold-safe feature-selection audit")
    ws.cell(detail_start - 1, 1).font = Font(color=NAVY, bold=True, size=12)
    _write_table_block(
        ws,
        frame,
        start_row=detail_start,
        table_name="T_FeatureSelectionDetail",
    )
    ws.freeze_panes = "A6"
    ws.print_title_rows = "1:5"
    ws.print_area = f"A1:{get_column_letter(summary.shape[1])}{summary_end}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return summary


def build_dl_experiment_sheet(ws, frame: pd.DataFrame) -> None:
    summary_columns = [
        "Run_ID", "Timestamp_UTC", "Environment", "GPU", "Dataset", "Seed",
        "Model_Variant", "Patch_Length", "Embedding_Dim", "KAN_Grid",
        "Parameter_Count", "Train_Time_Min", "Peak_VRAM_GB", "MAE_pp",
        "RMSE_pp", "Macro_F1", "Status",
    ]
    summary = frame.loc[:, [column for column in summary_columns if column in frame]].head(12)
    if summary.shape[1] == 0:
        summary = pd.DataFrame({"Status": ["No deep-learning runs recorded"]})
    _style_title(
        ws,
        summary.shape[1],
        "Deep-learning experiment ledger",
        "Compact compute/performance summary for visual review; the exact 33-column experiment ledger follows below.",
    )
    ws["A4"] = "Visible run summary"
    ws["A4"].font = Font(color=NAVY, bold=True, size=12)
    _, summary_end = _write_table_block(
        ws, summary, start_row=5, table_name="T_DL_Summary", max_rows=12
    )
    detail_start = max(22, summary_end + 4)
    ws.cell(detail_start - 1, 1, "Complete experiment ledger")
    ws.cell(detail_start - 1, 1).font = Font(color=NAVY, bold=True, size=12)
    _write_table_block(
        ws,
        frame,
        start_row=detail_start,
        table_name="T_13_DL_EXPERIMENTS",
    )
    ws.freeze_panes = "A6"
    ws.print_title_rows = "1:5"
    ws.print_area = f"A1:{get_column_letter(summary.shape[1])}{summary_end}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _benchmark_summary(benchmarks: pd.DataFrame, limit: int = 12) -> pd.DataFrame:
    if benchmarks.empty:
        return pd.DataFrame({"Status": ["No benchmark results recorded"]})
    work = benchmarks.copy()
    if "Status" in work:
        ok = work["Status"].astype(str).str.casefold().isin({"ok", "passed", "complete"})
        if ok.any():
            work = work.loc[ok]
    keys = [column for column in ("Task", "Model", "Feature_Set", "split_role") if column in work]
    metrics = [
        column for column in ("MAE_pp", "RMSE_pp", "R2", "Macro_F1", "Balanced_Accuracy")
        if column in work
    ]
    for column in metrics:
        work[column] = pd.to_numeric(work[column], errors="coerce")
    if keys and metrics:
        summary = work.groupby(keys, dropna=False, as_index=False)[metrics].mean()
    else:
        keep = keys + metrics + (["Status"] if "Status" in work else [])
        summary = work.loc[:, keep] if keep else pd.DataFrame({"Status": ["Benchmark schema unavailable"]})
    if "split_role" in summary:
        split_priority = {
            "final_holdout": 0,
            "development_validation": 1,
            "external_validation": 2,
        }
        summary["_split_priority"] = summary["split_role"].map(split_priority).fillna(9)
    else:
        summary["_split_priority"] = 9
    sort_columns = [column for column in ("_split_priority", "Task", "MAE_pp") if column in summary]
    if sort_columns:
        summary = summary.sort_values(sort_columns, ascending=True)
    summary = summary.drop(columns="_split_priority", errors="ignore").head(limit).reset_index(drop=True)
    summary["Status"] = "observed"
    return summary


def build_benchmarks_ablation_sheet(
    ws, benchmarks: pd.DataFrame, registry: pd.DataFrame
) -> pd.DataFrame:
    summary = _benchmark_summary(benchmarks)
    _style_title(
        ws,
        max(summary.shape[1], registry.shape[1]),
        "Benchmarks and ablations",
        "Observed comparison summary and pre-registered ablations are separate auditable tables; complete benchmark rows follow below.",
    )
    ws["A4"] = "Observed model comparison"
    ws["A4"].font = Font(color=NAVY, bold=True, size=12)
    _, summary_end = _write_table_block(
        ws, summary, start_row=5, table_name="T_BenchmarkSummary", max_rows=12
    )
    registry_label_row = summary_end + 3
    registry_start = registry_label_row + 1
    ws.cell(registry_label_row, 1, "Pre-registered ablation matrix")
    ws.cell(registry_label_row, 1).font = Font(color=NAVY, bold=True, size=12)
    _, registry_end = _write_table_block(
        ws, registry, start_row=registry_start, table_name="T_AblationRegistry"
    )
    detail_label_row = registry_end + 3
    detail_start = detail_label_row + 1
    ws.cell(detail_label_row, 1, "Complete observed benchmark ledger")
    ws.cell(detail_label_row, 1).font = Font(color=NAVY, bold=True, size=12)
    _write_table_block(
        ws,
        benchmarks,
        start_row=detail_start,
        table_name="T_BenchmarkDetails",
    )
    ws.freeze_panes = "A6"
    ws.print_title_rows = "1:5"
    preview_width = max(summary.shape[1], registry.shape[1])
    ws.print_area = f"A1:{get_column_letter(preview_width)}{registry_end}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return summary


def _behavior_diagnostics_from_summary(summary: dict[str, Any]) -> pd.DataFrame:
    behavior = summary.get("behavior", {}) if isinstance(summary, dict) else {}
    rows: list[dict[str, Any]] = []
    for fold in behavior.get("folds", []) or []:
        ari = pd.to_numeric(pd.Series([fold.get("bootstrap_ari")]), errors="coerce").iloc[0]
        exploratory = bool(fold.get("exploratory", not (pd.notna(ari) and ari >= 0.75)))
        rows.append(
            {
                "Diagnostic_Type": "Fold-local cluster selection",
                "Fold": fold.get("fold"),
                "Method": "GMM",
                "Candidate_K": "2–6",
                "Selected_K": fold.get("selected_k"),
                "Bootstrap_ARI": ari,
                "Selected_Status": "selected",
                "Stability_Status": "exploratory" if exploratory else "stable",
                "Notes": (
                    f"train_groups={fold.get('train_groups', 'NA')}; "
                    f"validation_groups={fold.get('validation_groups', 'NA')}"
                ),
                "Diagnostics_Source": "analysis_summary.json fallback",
            }
        )
    if behavior:
        final_ari = pd.to_numeric(
            pd.Series([behavior.get("final_cluster_bootstrap_ari")]), errors="coerce"
        ).iloc[0]
        final_exploratory = bool(
            behavior.get(
                "final_cluster_exploratory",
                not (pd.notna(final_ari) and final_ari >= 0.75),
            )
        )
        rows.append(
            {
                "Diagnostic_Type": "Final development clustering",
                "Fold": "development",
                "Method": "GMM",
                "Candidate_K": "2–6",
                "Selected_K": behavior.get("final_cluster_k"),
                "Bootstrap_ARI": final_ari,
                "Selected_Status": "selected",
                "Stability_Status": "exploratory" if final_exploratory else "stable",
                "Notes": "Final development-only cluster fit; untouched trips excluded",
                "Diagnostics_Source": "analysis_summary.json fallback",
            }
        )
    else:
        rows.append(
            {
                "Diagnostic_Type": "Cluster selection",
                "Fold": "pending",
                "Method": "GMM",
                "Candidate_K": "2–6",
                "Selected_K": None,
                "Bootstrap_ARI": None,
                "Selected_Status": "pending",
                "Stability_Status": "not reported",
                "Notes": "No diagnostics CSV or behaviour section was available",
                "Diagnostics_Source": "analysis_summary.json fallback",
            }
        )
    for method in ("KMeans", "HDBSCAN"):
        rows.append(
            {
                "Diagnostic_Type": "Comparator",
                "Fold": "development",
                "Method": method,
                "Candidate_K": "not persisted",
                "Selected_K": None,
                "Bootstrap_ARI": None,
                "Selected_Status": "comparison unavailable",
                "Stability_Status": "not reported",
                "Notes": "Detailed comparator diagnostics were not present in analysis_summary.json",
                "Diagnostics_Source": "analysis_summary.json fallback",
            }
        )
    for target, key in (
        ("Behaviour classifier", "selected_behavior_classifier"),
        ("Driver-ID diagnostic", "selected_driver_classifier"),
    ):
        if behavior.get(key):
            rows.append(
                {
                    "Diagnostic_Type": target,
                    "Fold": "development selection",
                    "Method": behavior[key],
                    "Candidate_K": None,
                    "Selected_K": None,
                    "Bootstrap_ARI": None,
                    "Selected_Status": "selected",
                    "Stability_Status": "not applicable",
                    "Notes": "Classifier selected on development validation only",
                    "Diagnostics_Source": "analysis_summary.json fallback",
                }
            )
    if not rows:
        rows = [
            {
                "Diagnostic_Type": "Cluster diagnostics",
                "Fold": "pending",
                "Method": method,
                "Candidate_K": "2–6" if method in {"GMM", "KMeans"} else "density-based",
                "Selected_K": None,
                "Bootstrap_ARI": None,
                "Selected_Status": "pending",
                "Stability_Status": "not reported",
                "Notes": "No diagnostics CSV or behaviour section was available",
                "Diagnostics_Source": "analysis_summary.json fallback",
            }
            for method in ("GMM", "KMeans", "HDBSCAN")
        ]
    return pd.DataFrame(rows)


def _prepare_behavior_diagnostics(
    diagnostics: pd.DataFrame, analysis_summary: dict[str, Any]
) -> pd.DataFrame:
    if diagnostics.empty:
        return _behavior_diagnostics_from_summary(analysis_summary)
    output = diagnostics.copy()
    selected_column = next(
        (
            column
            for column in ("Selected", "Is_Selected", "is_selected", "selected")
            if column in output
        ),
        None,
    )
    if "Selected_Status" not in output:
        if selected_column:
            output["Selected_Status"] = np.where(
                output[selected_column].fillna(False).astype(bool), "selected", "candidate"
            )
        else:
            output["Selected_Status"] = "reported"
    ari_column = next(
        (
            column
            for column in ("Bootstrap_ARI", "bootstrap_ari", "ARI", "ari")
            if column in output
        ),
        None,
    )
    if "Stability_Status" not in output:
        if ari_column:
            ari = pd.to_numeric(output[ari_column], errors="coerce")
            output["Stability_Status"] = np.where(
                ari.isna(), "not reported", np.where(ari >= 0.75, "stable", "exploratory")
            )
        else:
            output["Stability_Status"] = "not reported"
    output["Diagnostics_Source"] = "behavior_cluster_diagnostics.csv"
    return _prioritize_columns(
        output,
        [
            "Diagnostic_Type", "Fold", "Method", "Algorithm", "Candidate_K", "Selected_K",
            "BIC", "Silhouette", "Davies_Bouldin", "Bootstrap_ARI", "Selected_Status",
            "Stability_Status", "Notes", "Diagnostics_Source",
        ],
    )


def _behavior_profile(classifications: pd.DataFrame) -> pd.DataFrame:
    if classifications.empty or "behavior_level" not in classifications:
        return pd.DataFrame({"Status": ["Behaviour profiles unavailable"]})
    metrics = {
        "speed_mean_kmh": "Mean_Speed_kmh",
        "speed_p95_kmh": "P95_Speed_kmh",
        "accel_pos_rms": "Accel_RMS",
        "decel_abs_rms": "Decel_RMS",
        "jerk_rms": "Jerk_RMS",
        "pedal_mean": "Pedal_Mean",
        "brake_active_fraction": "Brake_Active_Fraction",
        "power_dis_p95_kw": "Power_Discharge_P95_kW",
        "bsi_intensity": "Mean_DBSI",
    }
    working = classifications.copy()
    for column in metrics:
        if column in working:
            working[column] = pd.to_numeric(working[column], errors="coerce")
    rows: list[dict[str, Any]] = []
    for level, subset in working.groupby("behavior_level", dropna=False):
        record: dict[str, Any] = {
            "Behavior_Level": level,
            "Trip_Count": int(len(subset)),
        }
        for source, destination in metrics.items():
            if source in subset:
                record[destination] = subset[source].mean()
        rows.append(record)
    profile = pd.DataFrame(rows)
    if profile.empty:
        return pd.DataFrame({"Status": ["Behaviour profiles unavailable"]})
    order = {"Eco": 0, "Balanced": 1, "Aggressive": 2}
    profile["_order"] = profile["Behavior_Level"].map(order).fillna(9)
    return profile.sort_values("_order").drop(columns="_order").reset_index(drop=True)


def build_behavior_sheet(
    ws,
    classifications: pd.DataFrame,
    diagnostics: pd.DataFrame,
) -> None:
    profile = _behavior_profile(classifications)
    _style_title(
        ws,
        max(diagnostics.shape[1], profile.shape[1]),
        "Driver-behaviour discovery",
        "Cluster candidates, comparator/stability evidence, behaviour profiles, and per-trip classifications are kept together.",
    )
    ws["A4"] = "Cluster selection and stability diagnostics"
    ws["A4"].font = Font(color=NAVY, bold=True, size=12)
    _, diagnostics_end = _write_table_block(
        ws,
        diagnostics,
        start_row=5,
        table_name="T_BehaviorDiagnostics",
        max_rows=40,
    )
    profile_label_row = diagnostics_end + 3
    profile_start = profile_label_row + 1
    ws.cell(profile_label_row, 1, "Cluster / behaviour profiles")
    ws.cell(profile_label_row, 1).font = Font(color=NAVY, bold=True, size=12)
    _, profile_end = _write_table_block(
        ws,
        profile,
        start_row=profile_start,
        table_name="T_BehaviorProfiles",
    )
    detail_label_row = profile_end + 3
    detail_start = detail_label_row + 1
    ws.cell(detail_label_row, 1, "Complete per-trip behaviour classifications")
    ws.cell(detail_label_row, 1).font = Font(color=NAVY, bold=True, size=12)
    _write_table_block(
        ws,
        classifications,
        start_row=detail_start,
        table_name="T_10_BEHAVIOR_CLUSTERS",
    )
    ws.freeze_panes = "A6"
    ws.print_title_rows = "1:5"
    preview_width = max(diagnostics.shape[1], profile.shape[1])
    ws.print_area = f"A1:{get_column_letter(preview_width)}{profile_end}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def build_external_validation() -> pd.DataFrame:
    return pd.DataFrame(
        [
            ["Nissan Leaf real-world", "Within-vehicle development", "~4 Hz / 1 Hz adapter", "BMS-reported SOH", "Expanding chronological trips", "Implemented", "Available", "Proof-of-concept only", "One vehicle; 127 trips; final 20% untouched"],
            ["RWTH EV/battery fleet", "External vehicle validation", "Adapter-defined", "Capacity-based SOH", "Leave-one-vehicle-out", "Implemented interface", "Not supplied", "Pending", "Provide source files locally; separate calibration head"],
            ["300-EV field dataset", "Low-cadence representation learning", "0.1 Hz native adapter", "Dataset-dependent", "Vehicle-held-out", "Implemented interface", "Not supplied", "Pending", "Do not upsample by repeated 1 Hz records"],
        ],
        columns=["Dataset", "Role", "Cadence", "SOH_Definition", "Vehicle_Split", "Adapter_Status", "Data_Status", "Claim_Status", "Notes"],
    )


def build_ablation_registry() -> pd.DataFrame:
    planned = {
        "Input modality": ["Battery only", "Driver only", "Battery + driver"],
        "Driver coupling": ["Concatenation", "One-way attention", "Bidirectional cross-attention"],
        "Pretraining": ["None", "Masked reconstruction", "Masked + alignment"],
        "Concepts": ["No bottleneck", "Concepts + MLP", "Concepts + KAN"],
        "D-BSI": ["None", "Total only", "Six components", "Components + cumulative dose"],
        "Reduction": ["All features", "Stability-selected", "PCA", "Sparse gates"],
        "Physics": ["No constraint", "Future-loss constraint", "Full soft consistency"],
        "KAN": ["MLP", "Spline KAN", "FastKAN", "Full FFN→KAN"],
        "Explanation": ["Latent residual off", "Latent residual on"],
        "Personalisation": ["Driver_ID excluded", "Driver_ID included diagnostic"],
    }
    return pd.DataFrame(
        [[family, variant, "Planned", "Must use identical outer folds and paired confidence intervals"]
         for family, variants in planned.items() for variant in variants],
        columns=["Ablation_Family", "Variant", "Status", "Decision_Rule"],
    )


def _find_column(frame: pd.DataFrame, choices: Iterable[str]) -> str | None:
    exact = {str(c).casefold(): str(c) for c in frame.columns}
    for choice in choices:
        if choice.casefold() in exact:
            return exact[choice.casefold()]
    return None


def build_dashboard(ws, tables: dict[str, pd.DataFrame]) -> None:
    _style_title(ws, 17, "DB-PatchKAN research dashboard", "Formula-linked overview. Metrics are descriptive unless their split is explicitly held out.")
    trip = tables["07_TRIP_FEATURES"]
    bsi = tables["11_BSI_RESULTS"]
    feature = tables["09_FEATURE_SELECTION"]
    benchmark = tables["14_BENCHMARKS_ABLATIONS"]

    cards = [
        ("Trips", len(trip), "Independent grouped unit"),
        ("Windows", len(tables["06_WINDOW_FEATURES"]), "Representation samples, not SOH labels"),
        ("Health pairs", len(tables["08_HEALTH_WINDOWS"]), "Forecast checkpoints"),
        ("Mean D-BSI", pd.to_numeric(bsi.get("bsi_intensity"), errors="coerce").mean() if not bsi.empty else np.nan, "0–100 stress intensity"),
        ("Cumulative dose", pd.to_numeric(bsi.get("stress_dose"), errors="coerce").sum() if not bsi.empty else np.nan, "Stress-weighted driving hours"),
        ("External vehicles", 0, "Validation pending: no fleet files supplied"),
    ]
    for index, (label, value, note) in enumerate(cards):
        row = 4 + (index // 3) * 4
        col = 1 + (index % 3) * 3
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 1)
        ws.cell(row, col, label)
        ws.cell(row, col).fill = PatternFill("solid", fgColor=TEAL)
        ws.cell(row, col).font = Font(color=WHITE, bold=True)
        ws.merge_cells(start_row=row + 1, start_column=col, end_row=row + 1, end_column=col + 1)
        ws.cell(row + 1, col, _safe_value(value))
        ws.cell(row + 1, col).font = Font(color=NAVY, bold=True, size=17)
        ws.merge_cells(start_row=row + 2, start_column=col, end_row=row + 2, end_column=col + 1)
        ws.cell(row + 2, col, note)
        ws.cell(row + 2, col).font = Font(color=SLATE, italic=True, size=9)

    ws["A13"] = "Trip"
    ws["B13"] = "SOH"
    ws["C13"] = "D-BSI"
    _format_header(ws[13][:3])
    trip_col = _find_column(trip, ["Trip", "trip_id", "trip_key"])
    soh_col = _find_column(trip, ["SOH", "soh_pct", "soh_target_raw_pct", "soh_mean_pct"])
    bsi_col = _find_column(trip, ["bsi_intensity"])
    n = min(len(trip), 127)
    for i in range(n):
        target_row = 14 + i
        source_row = 5 + i
        for col, source_name in enumerate((trip_col, soh_col, bsi_col), 1):
            if source_name:
                source_col = trip.columns.get_loc(source_name) + 1
                ws.cell(target_row, col, f"='07_TRIP_FEATURES'!{get_column_letter(source_col)}{source_row}")

    if n:
        soh_chart = LineChart()
        soh_chart.title = "SOH by chronological trip"
        soh_chart.y_axis.title = "SOH (%)"
        soh_chart.x_axis.title = "Trip"
        soh_chart.style = 13
        soh_chart.height = 7.2
        soh_chart.width = 14.5
        soh_chart.varyColors = False
        soh_chart.legend = None
        soh_chart.add_data(Reference(ws, min_col=2, min_row=13, max_row=13 + n), titles_from_data=True)
        soh_chart.set_categories(Reference(ws, min_col=1, min_row=14, max_row=13 + n))
        ws.add_chart(soh_chart, "J3")

        bsi_chart = LineChart()
        bsi_chart.title = "D-BSI intensity by trip"
        bsi_chart.y_axis.title = "D-BSI (0–100)"
        bsi_chart.x_axis.title = "Trip"
        bsi_chart.style = 12
        bsi_chart.height = 7.2
        bsi_chart.width = 14.5
        bsi_chart.varyColors = False
        bsi_chart.legend = None
        bsi_chart.add_data(Reference(ws, min_col=3, min_row=13, max_row=13 + n), titles_from_data=True)
        bsi_chart.set_categories(Reference(ws, min_col=1, min_row=14, max_row=13 + n))
        ws.add_chart(bsi_chart, "J18")

    ws["E13"] = "D-BSI component"
    ws["F13"] = "Mean contribution"
    _format_header(ws[13][4:6])
    components = ["rate", "thermal", "soc", "cycle", "transient", "interaction"]
    for i, name in enumerate(components, 14):
        ws.cell(i, 5, name.title())
        contribution = f"bsi_contribution_{name}"
        if contribution in bsi:
            col = bsi.columns.get_loc(contribution) + 1
            end = 4 + len(bsi)
            ws.cell(i, 6, f"=AVERAGE('11_BSI_RESULTS'!{get_column_letter(col)}5:{get_column_letter(col)}{end})")
    contribution_chart = BarChart()
    contribution_chart.type = "bar"
    contribution_chart.title = "Mean weighted D-BSI contribution"
    contribution_chart.y_axis.title = "Component"
    contribution_chart.x_axis.title = "D-BSI points"
    contribution_chart.style = 10
    contribution_chart.height = 7.2
    contribution_chart.width = 14.5
    contribution_chart.add_data(Reference(ws, min_col=6, min_row=13, max_row=19), titles_from_data=True)
    contribution_chart.set_categories(Reference(ws, min_col=5, min_row=14, max_row=19))
    contribution_chart.dLbls = DataLabelList()
    contribution_chart.dLbls.showVal = True
    ws.add_chart(contribution_chart, "J33")

    selected = feature.copy()
    importance_col = _find_column(selected, ["Permutation_Importance", "permutation_importance", "Consensus_Rank"])
    name_col = _find_column(selected, ["Feature_Name", "feature"])
    if name_col and importance_col:
        selected[importance_col] = pd.to_numeric(selected[importance_col], errors="coerce")
        selected = selected.nlargest(10, importance_col)
        ws["E22"] = "Top feature"
        ws["F22"] = "Held-out importance"
        _format_header(ws[22][4:6])
        for row, (_, record) in enumerate(selected.iterrows(), 23):
            ws.cell(row, 5, _safe_value(record[name_col]))
            ws.cell(row, 6, _safe_value(record[importance_col]))

    ws["A145"] = "Decision rule"
    ws["B145"] = "Recommend DB-PatchKAN only when paired external-vehicle confidence intervals beat both operational XGBoost/EBM and the parameter-matched Patch Transformer + MLP."
    ws.merge_cells("B145:H147")
    ws["A145"].font = Font(color=WHITE, bold=True)
    ws["A145"].fill = PatternFill("solid", fgColor=NAVY)
    ws["B145"].alignment = Alignment(wrap_text=True, vertical="top")
    for col in range(1, 18):
        ws.column_dimensions[get_column_letter(col)].width = 14
    ws.freeze_panes = "A13"
    ws.print_area = "A1:X48"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1


def build_dashboard_v2(ws, tables: dict[str, pd.DataFrame]) -> None:
    """Build the formula-linked dashboard required by the research protocol."""

    _style_title(
        ws,
        24,
        "DB-PatchKAN research dashboard",
        "Formula-linked overview. Metrics are descriptive unless their split is explicitly held out.",
    )
    trip = tables["07_TRIP_FEATURES"]
    bsi = tables["11_BSI_RESULTS"]
    behavior = tables["10_BEHAVIOR_CLUSTERS"]
    feature_summary = tables.get("_FEATURE_SUMMARY", pd.DataFrame())
    benchmark_summary = tables.get("_BENCHMARK_SUMMARY", pd.DataFrame())

    cards = [
        (
            "Trips",
            "=ROWS(T_07_TRIP_FEATURES[trip_key])" if "trip_key" in trip else "=0",
            "Independent grouped unit",
        ),
        (
            "Windows",
            "=ROWS(T_06_WINDOW_FEATURES[window_id])"
            if "window_id" in tables["06_WINDOW_FEATURES"]
            else "=0",
            "Representation samples, not SOH labels",
        ),
        (
            "Health pairs",
            "=ROWS(T_08_HEALTH_WINDOWS[health_window_id])"
            if "health_window_id" in tables["08_HEALTH_WINDOWS"]
            else "=0",
            "Forecast checkpoints",
        ),
        (
            "Mean D-BSI",
            "=AVERAGE(T_11_BSI_RESULTS[bsi_intensity])" if "bsi_intensity" in bsi else "=0",
            "0–100 stress intensity",
        ),
        (
            "Cumulative dose",
            "=SUM(T_11_BSI_RESULTS[stress_dose])" if "stress_dose" in bsi else "=0",
            "Stress-weighted driving hours",
        ),
        (
            "External vehicles",
            '=COUNTIF(T_16_EXTERNAL_VALIDATION[Claim_Status],"Validated")',
            "Validation pending until fleet files are supplied",
        ),
    ]
    for index, (label, formula, note) in enumerate(cards):
        row = 4 + (index // 3) * 4
        col = 1 + (index % 3) * 3
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 1)
        ws.cell(row, col, label)
        ws.cell(row, col).fill = PatternFill("solid", fgColor=TEAL)
        ws.cell(row, col).font = Font(color=WHITE, bold=True)
        ws.merge_cells(start_row=row + 1, start_column=col, end_row=row + 1, end_column=col + 1)
        ws.cell(row + 1, col, formula)
        ws.cell(row + 1, col).font = Font(color=NAVY, bold=True, size=17)
        ws.cell(row + 1, col).number_format = (
            "0.000" if label in {"Mean D-BSI", "Cumulative dose"} else "0"
        )
        ws.merge_cells(start_row=row + 2, start_column=col, end_row=row + 2, end_column=col + 1)
        ws.cell(row + 2, col, note)
        ws.cell(row + 2, col).font = Font(color=SLATE, italic=True, size=9)

    ws["A13"] = "Trip"
    ws["B13"] = "SOH"
    ws["C13"] = "D-BSI"
    _format_header(ws[13][:3])
    trip_col = _find_column(trip, ["Trip", "trip_id", "trip_key"])
    soh_col = _find_column(trip, ["SOH", "soh_pct", "soh_target_raw_pct", "soh_mean_pct"])
    bsi_col = _find_column(trip, ["bsi_intensity"])
    n = min(len(trip), 127)
    for index in range(n):
        target_row = 14 + index
        source_row = 5 + index
        for col, source_name in enumerate((trip_col, soh_col, bsi_col), 1):
            if source_name:
                source_col = trip.columns.get_loc(source_name) + 1
                ws.cell(
                    target_row,
                    col,
                    f"='07_TRIP_FEATURES'!{get_column_letter(source_col)}{source_row}",
                )

    if n:
        soh_chart = LineChart()
        soh_chart.title = "SOH by chronological trip"
        soh_chart.y_axis.title = "SOH (%)"
        soh_chart.x_axis.title = "Trip"
        soh_chart.style = 13
        soh_chart.height = 7.2
        soh_chart.width = 14.5
        soh_chart.varyColors = False
        soh_chart.legend = None
        soh_chart.add_data(
            Reference(ws, min_col=2, min_row=13, max_row=13 + n), titles_from_data=True
        )
        soh_chart.set_categories(Reference(ws, min_col=1, min_row=14, max_row=13 + n))
        ws.add_chart(soh_chart, "J3")

        bsi_chart = LineChart()
        bsi_chart.title = "D-BSI intensity by trip"
        bsi_chart.y_axis.title = "D-BSI (0–100)"
        bsi_chart.x_axis.title = "Trip"
        bsi_chart.style = 12
        bsi_chart.height = 7.2
        bsi_chart.width = 14.5
        bsi_chart.varyColors = False
        bsi_chart.legend = None
        bsi_chart.add_data(
            Reference(ws, min_col=3, min_row=13, max_row=13 + n), titles_from_data=True
        )
        bsi_chart.set_categories(Reference(ws, min_col=1, min_row=14, max_row=13 + n))
        ws.add_chart(bsi_chart, "J18")

    ws["E13"] = "Behaviour level"
    ws["F13"] = "Trip count"
    ws["G13"] = "Proportion"
    _format_header(ws[13][4:7])
    for row, level in enumerate(("Eco", "Balanced", "Aggressive"), 14):
        ws.cell(row, 5, level)
        if "behavior_level" in behavior:
            ws.cell(row, 6, f'=COUNTIF(T_10_BEHAVIOR_CLUSTERS[behavior_level],E{row})')
            ws.cell(row, 7, f"=IFERROR(F{row}/SUM($F$14:$F$16),0)")
        else:
            ws.cell(row, 6, 0)
            ws.cell(row, 7, 0)
        ws.cell(row, 7).number_format = "0.0%"

    ws["E20"] = "D-BSI component"
    ws["F20"] = "Mean contribution"
    _format_header(ws[20][4:6])
    for row, name in enumerate(
        ("rate", "thermal", "soc", "cycle", "transient", "interaction"), 21
    ):
        ws.cell(row, 5, name.title())
        contribution = f"bsi_contribution_{name}"
        if contribution in bsi:
            source_col = bsi.columns.get_loc(contribution) + 1
            source_end = 4 + len(bsi)
            ws.cell(
                row,
                6,
                f"=AVERAGE('11_BSI_RESULTS'!{get_column_letter(source_col)}5:"
                f"{get_column_letter(source_col)}{source_end})",
            )

    ws["E29"] = "Top consensus feature"
    ws["F29"] = "Held-out importance"
    _format_header(ws[29][4:6])
    feature_name_col = _find_column(feature_summary, ["Feature_Name", "feature"])
    feature_importance_col = _find_column(
        feature_summary, ["Permutation_Importance", "Importance"]
    )
    if feature_name_col and feature_importance_col:
        name_col = feature_summary.columns.get_loc(feature_name_col) + 1
        importance_col = feature_summary.columns.get_loc(feature_importance_col) + 1
        for offset in range(min(8, len(feature_summary))):
            source_row = 6 + offset
            ws.cell(
                30 + offset,
                5,
                f"='09_FEATURE_SELECTION'!{get_column_letter(name_col)}{source_row}",
            )
            ws.cell(
                30 + offset,
                6,
                f"='09_FEATURE_SELECTION'!{get_column_letter(importance_col)}{source_row}",
            )

    model_columns = [
        column
        for column in ("Task", "Model", "split_role", "MAE_pp", "Macro_F1")
        if column in benchmark_summary
    ]
    if model_columns:
        for col, value in enumerate(model_columns, 5):
            ws.cell(40, col, value)
        _format_header(ws[40][4 : 4 + len(model_columns)])
        for offset in range(min(6, len(benchmark_summary))):
            for col, source_name in enumerate(model_columns, 5):
                source_col = benchmark_summary.columns.get_loc(source_name) + 1
                ws.cell(
                    41 + offset,
                    col,
                    f"='14_BENCHMARKS_ABLATIONS'!{get_column_letter(source_col)}{6 + offset}",
                )
    else:
        ws["E40"] = "Model comparison"
        ws["F40"] = "Pending"
        _format_header(ws[40][4:6])

    behavior_chart = BarChart()
    behavior_chart.type = "col"
    behavior_chart.title = "Discovered behaviour mix"
    behavior_chart.y_axis.title = "Trips"
    behavior_chart.x_axis.title = "Behaviour level"
    behavior_chart.style = 10
    behavior_chart.height = 6.2
    behavior_chart.width = 14.5
    behavior_chart.varyColors = True
    behavior_chart.legend = None
    behavior_chart.add_data(
        Reference(ws, min_col=6, min_row=13, max_row=16), titles_from_data=True
    )
    behavior_chart.set_categories(Reference(ws, min_col=5, min_row=14, max_row=16))
    behavior_chart.dLbls = DataLabelList()
    behavior_chart.dLbls.showLegendKey = False
    behavior_chart.dLbls.showVal = True
    behavior_chart.dLbls.showCatName = False
    behavior_chart.dLbls.showSerName = False
    behavior_chart.dLbls.showPercent = False
    behavior_chart.dLbls.showBubbleSize = False
    ws.add_chart(behavior_chart, "J33")

    # Keep the decision note outside A:C, which contain the complete 127-trip
    # chart source through row 140.  Placing it in A49 previously overwrote
    # three trip records and introduced a false SOH discontinuity.
    ws["E48"] = "Decision rule"
    ws["F48"] = (
        "Recommend DB-PatchKAN only when paired external-vehicle confidence intervals beat "
        "both operational XGBoost/EBM and the parameter-matched Patch Transformer + MLP."
    )
    ws.merge_cells("F48:I51")
    ws["E48"].font = Font(color=WHITE, bold=True)
    ws["E48"].fill = PatternFill("solid", fgColor=NAVY)
    ws["F48"].alignment = Alignment(wrap_text=True, vertical="top")
    for col in range(1, 25):
        ws.column_dimensions[get_column_letter(col)].width = 14
    ws.freeze_panes = "A13"
    ws.print_area = "A1:X52"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def build_bsi_audit(ws, bsi: pd.DataFrame) -> None:
    title = "Formula-driven D-BSI audit"
    subtitle = "Editable component weights are linked from 01_PARAMETERS. Audit differences must be zero within numeric precision."
    columns = [
        "Trip_Key", "Duration_Hours", "Rate", "Thermal", "SOC", "Cycle",
        "Transient", "Interaction", "Rate_Contribution", "Thermal_Contribution",
        "SOC_Contribution", "Cycle_Contribution", "Transient_Contribution",
        "Interaction_Contribution", "Audit_DBSI", "Pipeline_DBSI", "DBSI_Difference",
        "Audit_Stress_Dose", "Pipeline_Stress_Dose", "Dose_Difference", "Audit_Status",
    ]
    _style_title(ws, len(columns), title, subtitle)
    for col, value in enumerate(columns, 1):
        ws.cell(4, col, value)
    _format_header(ws[4])
    key_col = _find_column(bsi, ["trip_key", "Trip", "trip_id"])
    source_cols = {
        "Rate": "bsi_rate", "Thermal": "bsi_thermal", "SOC": "bsi_soc",
        "Cycle": "bsi_cycle", "Transient": "bsi_transient", "Interaction": "bsi_interaction",
    }
    required = [
        "duration_s", *source_cols.values(), "bsi_intensity", "stress_dose"
    ]
    missing_columns = [column for column in required if column not in bsi]
    if bsi.empty:
        raise ValueError("D-BSI audit requires at least one pipeline result row")
    if missing_columns:
        raise ValueError(f"D-BSI audit is missing required columns: {missing_columns}")
    invalid_counts = {
        column: int(pd.to_numeric(bsi[column], errors="coerce").isna().sum())
        for column in required
    }
    invalid_counts = {column: count for column, count in invalid_counts.items() if count}
    if invalid_counts:
        raise ValueError(f"D-BSI audit contains non-numeric or missing inputs: {invalid_counts}")
    n = len(bsi)
    for i in range(n):
        row = 5 + i
        if i < len(bsi):
            rec = bsi.iloc[i]
            ws.cell(row, 1, _safe_value(rec.get(key_col)) if key_col else i + 1)
            duration_s = pd.to_numeric(pd.Series([rec.get("duration_s")]), errors="coerce").iloc[0]
            ws.cell(row, 2, float(duration_s) / 3600.0 if pd.notna(duration_s) else None)
            for col, source in enumerate(source_cols.values(), 3):
                ws.cell(row, col, _safe_value(rec.get(source)))
            ws.cell(row, 16, _safe_value(rec.get("bsi_intensity")))
            ws.cell(row, 19, _safe_value(rec.get("stress_dose")))
        weights = [3, 4, 5, 6, 7, 8]
        for offset, weight_row in enumerate(weights, 9):
            component_col = get_column_letter(offset - 6)
            ws.cell(row, offset, f"={component_col}{row}*'01_PARAMETERS'!$I${weight_row}*100")
        ws.cell(row, 15, f"=SUM(I{row}:N{row})")
        ws.cell(row, 17, f"=O{row}-P{row}")
        ws.cell(row, 18, f"=O{row}/100*B{row}")
        ws.cell(row, 20, f"=R{row}-S{row}")
        ws.cell(row, 21, f'=IF(AND(ABS(Q{row})<0.000001,ABS(T{row})<0.000001),"PASS","CHECK")')
        for col in range(2, 21):
            ws.cell(row, col).number_format = "0.000000"
    end = 4 + n
    table = Table(displayName="T_BSI_Audit", ref=f"A4:U{end}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    ws.conditional_formatting.add(f"Q5:Q{end}", ColorScaleRule(start_type="min", start_color=GREEN, mid_type="percentile", mid_value=50, mid_color=AMBER, end_type="max", end_color=RED))
    ws.conditional_formatting.add(f"T5:T{end}", ColorScaleRule(start_type="min", start_color=GREEN, mid_type="percentile", mid_value=50, mid_color=AMBER, end_type="max", end_color=RED))
    ws.freeze_panes = "C5"
    for col in range(1, 22):
        ws.column_dimensions[get_column_letter(col)].width = 18 if col > 1 else 32
    ws.print_area = f"A1:U{min(end, 40)}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1


def build_workbook(input_dir: Path, output_path: Path, config_path: Path, source_path: Path) -> Path:
    config = yaml.safe_load(config_path.read_text(encoding="utf-8"))
    manifest_path = input_dir / "manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8")) if manifest_path.exists() else {}
    analysis_summary_path = input_dir / "analysis_summary.json"
    analysis_summary = (
        json.loads(analysis_summary_path.read_text(encoding="utf-8"))
        if analysis_summary_path.exists()
        else {}
    )

    tables: dict[str, pd.DataFrame] = {
        "02_DATA_DICTIONARY": _load_frame(input_dir, ["data_dictionary.csv"]),
        "03_SOURCE_PROFILE": _load_frame(input_dir, ["profile.csv", "source_profile.csv"]),
        "04_QC_SUMMARY": _load_frame(input_dir, ["qc_summary.csv"]),
        "05_RAW_SAMPLE": _load_frame(input_dir, ["raw_sample.csv"]),
        "06_WINDOW_FEATURES": _load_frame(input_dir, ["window_features.csv"]),
        "07_TRIP_FEATURES": _load_frame(input_dir, ["trip_features.csv"]),
        "08_HEALTH_WINDOWS": _load_frame(input_dir, ["health_windows.csv"]),
        "09_FEATURE_SELECTION": _load_frame(input_dir, ["feature_selection.csv"], columns=FEATURE_SELECTION_COLUMNS),
        "10_BEHAVIOR_CLUSTERS": _load_frame(input_dir, ["behavior_clusters.csv", "behaviour_clusters.csv"]),
        "11_BSI_RESULTS": _load_frame(input_dir, ["bsi_results.csv"]),
        "12_SOH_PREDICTIONS": _load_frame(input_dir, ["soh_predictions.csv"]),
        "13_DL_EXPERIMENTS": _load_frame(input_dir, ["deep_experiments.csv", "dl_experiments.csv"], columns=DL_EXPERIMENT_COLUMNS),
        "15_EXPLAINABILITY": _load_frame(input_dir, ["explainability.csv"]),
    }
    tables["05_RAW_SAMPLE"] = _stratified_raw_excerpt(tables["05_RAW_SAMPLE"])
    # Excel table headers are case-insensitive.  Preserve both the source and
    # cleaned channels while avoiding ECO/eco and ePedal/epedal collisions.
    tables["05_RAW_SAMPLE"] = tables["05_RAW_SAMPLE"].rename(
        columns={"ECO": "ECO_source_raw", "ePedal": "ePedal_source_raw"}
    )
    benchmarks = _load_frame(input_dir, ["benchmarks.csv", "benchmarks_ablations.csv"])
    behavior_diagnostics = _prepare_behavior_diagnostics(
        _load_frame(input_dir, ["behavior_cluster_diagnostics.csv"]),
        analysis_summary,
    )
    tables["14_BENCHMARKS_ABLATIONS"] = benchmarks
    tables["16_EXTERNAL_VALIDATION"] = build_external_validation()

    # Put the decision-bearing fields inside each sheet's visible preview while
    # retaining every remaining column in the same worksheet table.
    tables["08_HEALTH_WINDOWS"] = _prioritize_columns(
        tables["08_HEALTH_WINDOWS"],
        [
            "health_window_id", "trip_key", "Trip", "Driver_ID", "anchor_time_utc",
            "soh_target_raw_pct", "soh_target_trend_pct", "cumulative_distance_km",
            "bsi_intensity", "stress_dose", "target_soh_30d_pct", "target_loss_30d_pp",
            "target_soh_500km_pct", "target_loss_500km_pp",
            "target_loss_rate_pp_per_1000km", "target_quality_flag",
        ],
    )
    tables["10_BEHAVIOR_CLUSTERS"] = _prioritize_columns(
        tables["10_BEHAVIOR_CLUSTERS"],
        [
            "trip_key", "Trip", "Driver_ID", "trip_start_utc", "natural_cluster_id",
            "aggression_score", "behavior_level", "behavior_level_pred", "driver_id_pred",
            "split_role", "Fold",
        ],
    )
    tables["11_BSI_RESULTS"] = _prioritize_columns(
        tables["11_BSI_RESULTS"],
        [
            "trip_key", "Trip", "Driver_ID", "duration_s", "bsi_rate", "bsi_thermal",
            "bsi_soc", "bsi_cycle", "bsi_transient", "bsi_interaction", "bsi_intensity",
            "stress_dose", "cumulative_stress_dose", "bsi_band", "bsi_version",
        ],
    )
    tables["12_SOH_PREDICTIONS"] = _prioritize_columns(
        tables["12_SOH_PREDICTIONS"],
        [
            "Task", "Model", "Feature_Set", "Fold", "split_role", "Row_Index",
            "target_observed", "target_predicted", "soh_observed_pct", "soh_pred_pct",
            "soh_residual_pct", "Prediction_Lower", "Prediction_Upper", "Target_Unit",
        ],
    )
    tables["15_EXPLAINABILITY"] = _prioritize_columns(
        tables["15_EXPLAINABILITY"],
        ["Task", "Model", "Feature_Name", "Explanation_Type", "Importance", "Status"],
    )

    wb = Workbook()
    default = wb.active
    wb.remove(default)
    for name in SHEETS:
        wb.create_sheet(name)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.properties.title = "DB-PatchKAN research workbook"
    wb.properties.subject = "Driver behaviour, D-BSI, feature selection, and EV battery SOH"
    wb.properties.creator = "DB-PatchKAN reproducible pipeline"
    wb.properties.description = "Formula-audited research workbook; external validation status is explicit."

    build_readme(wb["00_README"], manifest, source_path, input_dir, config_path)
    build_parameters(wb["01_PARAMETERS"], config)
    for name in [
        sheet
        for sheet in SHEETS
        if sheet in SHEET_SPECS
        and sheet
        not in {
            "09_FEATURE_SELECTION",
            "10_BEHAVIOR_CLUSTERS",
            "13_DL_EXPERIMENTS",
            "14_BENCHMARKS_ABLATIONS",
        }
    ]:
        title, subtitle = SHEET_SPECS[name]
        write_frame(wb[name], tables[name], title, subtitle)
        if name in {"11_BSI_RESULTS", "12_SOH_PREDICTIONS"}:
            used = wb[name].max_row
            if used >= 5:
                numeric_columns = [
                    i + 1 for i, col in enumerate(tables[name].columns)
                    if any(token in str(col).casefold() for token in ("importance", "frequency", "probability", "bsi", "mae", "rmse", "r2", "residual"))
                ]
                for col in numeric_columns[:6]:
                    letter = get_column_letter(col)
                    wb[name].conditional_formatting.add(
                        f"{letter}5:{letter}{used}",
                        DataBarRule(start_type="min", end_type="max", color=BLUE),
                    )
    tables["_FEATURE_SUMMARY"] = build_feature_selection_sheet(
        wb["09_FEATURE_SELECTION"], tables["09_FEATURE_SELECTION"]
    )
    build_behavior_sheet(
        wb["10_BEHAVIOR_CLUSTERS"],
        tables["10_BEHAVIOR_CLUSTERS"],
        behavior_diagnostics,
    )
    build_dl_experiment_sheet(wb["13_DL_EXPERIMENTS"], tables["13_DL_EXPERIMENTS"])
    tables["_BENCHMARK_SUMMARY"] = build_benchmarks_ablation_sheet(
        wb["14_BENCHMARKS_ABLATIONS"], benchmarks, build_ablation_registry()
    )
    build_dashboard_v2(wb["17_DASHBOARD"], tables)
    build_bsi_audit(wb["18_BSI_AUDIT"], tables["11_BSI_RESULTS"])

    # Hyperlinks in the README turn the workbook into a navigable research artifact.
    readme = wb["00_README"]
    readme["D23"] = "Workbook navigation"
    readme["D23"].font = Font(color=NAVY, bold=True, size=13)
    for index, name in enumerate(SHEETS[1:], 24):
        readme.cell(index, 4, name)
        readme.cell(index, 4).hyperlink = f"#'{name}'!A1"
        readme.cell(index, 4).style = "Hyperlink"
    readme.print_area = f"A1:H{23 + len(SHEETS) - 1}"

    # Workbook-wide consistency pass.
    for ws in wb.worksheets:
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        ws.page_margins.left = 0.25
        ws.page_margins.right = 0.25
        ws.page_margins.top = 0.5
        ws.page_margins.bottom = 0.5
        ws.oddFooter.center.text = "DB-PatchKAN | leakage-safe research artifact"
        ws.oddFooter.right.text = "Page &P of &N"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--config", type=Path, default=Path("configs/default.yaml"))
    parser.add_argument("--source", type=Path, default=Path("Nissan Leaf Real-World Data (1) (1).csv"))
    args = parser.parse_args()
    result = build_workbook(
        args.input_dir.resolve(), args.output.resolve(), args.config.resolve(), args.source.resolve()
    )
    print(result)


if __name__ == "__main__":
    main()
