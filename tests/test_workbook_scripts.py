from __future__ import annotations

import json
import sys
from datetime import datetime
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries


ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

from build_workbook import (  # noqa: E402
    SHEETS,
    _prepare_behavior_diagnostics,
    build_workbook,
)
from verify_workbook import verify  # noqa: E402


def _write_minimal_inputs(directory: Path, *, include_bsi: bool = True) -> None:
    directory.mkdir(parents=True, exist_ok=True)
    (directory / "manifest.json").write_text(
        json.dumps(
            {
                "source_sha256": "abc123",
                "rows_processed": 2,
                "pipeline": "test-pipeline",
                "pipeline_patch_level": "test-patch",
                "qc_patch_level": "test-qc",
            }
        ),
        encoding="utf-8",
    )
    pd.DataFrame(
        {
            "trip_key": ["vehicle:1", "vehicle:2"],
            "Trip": [1, 2],
            "soh_target_raw_pct": [96.0, 95.9],
            "bsi_intensity": [10.0, 20.0],
        }
    ).to_csv(directory / "trip_features.csv", index=False)
    pd.DataFrame({"window_id": ["w1"], "trip_key": ["vehicle:1"]}).to_csv(
        directory / "window_features.csv", index=False
    )
    pd.DataFrame(
        {
            "health_window_id": ["h1"],
            "trip_key": ["vehicle:1"],
            "anchor_time_utc": ["2026-01-01T00:00:00+00:00"],
            "soh_target_raw_pct": [96.0],
            "target_soh_30d_pct": [95.9],
            "target_loss_30d_pp": [0.1],
            "target_soh_500km_pct": [95.8],
            "target_loss_500km_pp": [0.2],
        }
    ).to_csv(directory / "health_windows.csv", index=False)
    pd.DataFrame(
        {
            "trip_key": ["vehicle:1", "vehicle:2"],
            "behavior_level": ["Eco", "Balanced"],
            "natural_cluster_id": [0, 1],
            "aggression_score": [0.1, 0.5],
        }
    ).to_csv(directory / "behavior_clusters.csv", index=False)
    pd.DataFrame(
        {
            "Feature_Name": ["jerk_rms", "age_days"],
            "Feature_Group": ["driver", "aging"],
            "Permutation_Importance": [0.2, 0.1],
            "Fold_Selection_Frequency": [1.0, 0.25],
            "Consensus_Rank": [1, 2],
            "Final_Status": ["selected", "rejected"],
            "Stage": ["development_consensus", "development_consensus"],
            "Task": ["current_soh", "current_soh"],
        }
    ).to_csv(directory / "feature_selection.csv", index=False)
    pd.DataFrame(
        {
            "Task": ["current_soh", "current_soh"],
            "Model": ["XGBoost", "ElasticNet"],
            "Feature_Set": ["Consensus", "Consensus"],
            "split_role": ["final_holdout", "final_holdout"],
            "Status": ["ok", "ok"],
            "MAE_pp": [0.2, 0.4],
            "RMSE_pp": [0.25, 0.45],
            "R2": [0.8, 0.6],
        }
    ).to_csv(directory / "benchmarks.csv", index=False)
    pd.DataFrame(
        {
            "Run_ID": ["run-1"],
            "Timestamp_UTC": ["2026-01-01T12:34:56+00:00"],
            "Environment": ["test"],
            "GPU": ["CPU"],
            "Dataset": ["synthetic"],
            "Seed": [42],
            "Model_Variant": ["DB-PatchKAN"],
            "Status": ["passed"],
        }
    ).to_csv(directory / "deep_experiments.csv", index=False)
    if include_bsi:
        pd.DataFrame(
            {
                "trip_key": ["vehicle:1", "vehicle:2"],
                "Trip": [1, 2],
                "Driver_ID": [0, 1],
                "duration_s": [3600.0, 1800.0],
                "bsi_rate": [0.1, 0.2],
                "bsi_thermal": [0.2, 0.1],
                "bsi_soc": [0.3, 0.2],
                "bsi_cycle": [0.1, 0.1],
                "bsi_transient": [0.2, 0.1],
                "bsi_interaction": [0.1, 0.2],
                "bsi_intensity": [16.5, 15.0],
                "stress_dose": [0.165, 0.075],
                "cumulative_stress_dose": [0.165, 0.240],
            }
        ).to_csv(directory / "bsi_results.csv", index=False)


def test_workbook_contract_and_formula_linkage(tmp_path: Path) -> None:
    inputs = tmp_path / "inputs"
    _write_minimal_inputs(inputs)
    source = tmp_path / "source.csv"
    source.write_text("x\n1\n", encoding="utf-8")
    output = tmp_path / "research.xlsx"
    build_workbook(inputs, output, ROOT / "configs" / "default.yaml", source)

    workbook = load_workbook(output, data_only=False)
    assert workbook.sheetnames == SHEETS
    assert workbook["01_PARAMETERS"]["I3"].value.startswith("=$C$")
    assert workbook["17_DASHBOARD"]["A5"].value == "=ROWS(T_07_TRIP_FEATURES[trip_key])"
    assert "COUNTIF(T_10_BEHAVIOR_CLUSTERS" in workbook["17_DASHBOARD"]["F14"].value
    assert "09_FEATURE_SELECTION" in workbook["17_DASHBOARD"]["E30"].value
    assert "14_BENCHMARKS_ABLATIONS" in workbook["17_DASHBOARD"]["E41"].value
    assert workbook["17_DASHBOARD"]["E48"].value == "Decision rule"
    assert str(workbook["17_DASHBOARD"]["A14"].value).startswith("='07_TRIP_FEATURES'!")
    assert set(workbook["14_BENCHMARKS_ABLATIONS"].tables) == {
        "T_BenchmarkSummary",
        "T_AblationRegistry",
        "T_BenchmarkDetails",
    }
    assert set(workbook["10_BEHAVIOR_CLUSTERS"].tables) == {
        "T_BehaviorDiagnostics",
        "T_BehaviorProfiles",
        "T_10_BEHAVIOR_CLUSTERS",
    }
    diagnostic_table = workbook["10_BEHAVIOR_CLUSTERS"].tables["T_BehaviorDiagnostics"]
    diag_min_col, diag_min_row, diag_max_col, diag_max_row = range_boundaries(
        diagnostic_table.ref
    )
    diag_headers = {
        workbook["10_BEHAVIOR_CLUSTERS"].cell(diag_min_row, column).value: column
        for column in range(diag_min_col, diag_max_col + 1)
    }
    methods = {
        workbook["10_BEHAVIOR_CLUSTERS"].cell(row, diag_headers["Method"]).value
        for row in range(diag_min_row + 1, diag_max_row + 1)
    }
    assert {"GMM", "KMeans", "HDBSCAN"}.issubset(methods)
    assert all(sheet.print_area for sheet in workbook.worksheets)

    detail = workbook["13_DL_EXPERIMENTS"].tables["T_13_DL_EXPERIMENTS"]
    min_col, min_row, _, _ = range_boundaries(detail.ref)
    headers = {
        workbook["13_DL_EXPERIMENTS"].cell(min_row, column).value: column
        for column in range(min_col, min_col + 33)
    }
    timestamp = workbook["13_DL_EXPERIMENTS"].cell(min_row + 1, headers["Timestamp_UTC"]).value
    assert isinstance(timestamp, datetime)
    workbook.close()

    with pytest.raises(AssertionError, match="recalculated numeric cache"):
        verify(output)


def test_builder_rejects_empty_dbsi_audit(tmp_path: Path) -> None:
    inputs = tmp_path / "inputs"
    _write_minimal_inputs(inputs, include_bsi=False)
    source = tmp_path / "source.csv"
    source.write_text("x\n1\n", encoding="utf-8")
    with pytest.raises(ValueError, match="requires at least one"):
        build_workbook(
            inputs,
            tmp_path / "research.xlsx",
            ROOT / "configs" / "default.yaml",
            source,
        )


def test_behavior_diagnostics_csv_gets_explicit_selection_and_stability_status() -> None:
    diagnostics = pd.DataFrame(
        {
            "Method": ["GMM", "KMeans", "HDBSCAN"],
            "Candidate_K": [2, 3, None],
            "Bootstrap_ARI": [0.80, 0.40, None],
            "Selected": [True, False, False],
        }
    )
    prepared = _prepare_behavior_diagnostics(diagnostics, {})
    assert prepared["Diagnostics_Source"].eq("behavior_cluster_diagnostics.csv").all()
    assert prepared["Selected_Status"].tolist() == ["selected", "candidate", "candidate"]
    assert prepared["Stability_Status"].tolist() == [
        "stable",
        "exploratory",
        "not reported",
    ]
